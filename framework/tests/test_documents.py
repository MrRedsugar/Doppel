import hashlib

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from doppel.documents import WorkspaceDocuments


def make_workbook(root, rows=1000):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sales'
    sheet.append(['Region', 'Amount', 'Units'])
    sheet['A1'].font = Font(bold=True, color='FF0000')
    for i in range(rows):
        sheet.append(['East' if i % 2 else 'West', i + 1, i % 5 + 1])
        sheet.cell(i + 2, 2).number_format = '#,##0.00'
    book.save(root / 'source.xlsx')
    book.close()


def test_inspection_is_bounded_and_aggregate_independent(tmp_path):
    make_workbook(tmp_path)
    documents = WorkspaceDocuments(tmp_path)
    before = hashlib.sha256((tmp_path / 'source.xlsx').read_bytes()).hexdigest()
    preview = documents.inspect_workbook('source.xlsx')
    assert preview['sheets'][0]['headers'] == ['Region', 'Amount', 'Units']
    assert len(preview['sheets'][0]['samples']) <= 5
    assert preview['sheets'][0]['rows'] == 1000
    result = documents.transform_workbook('source.xlsx', 'result.xlsx', [
        {'op': 'group_sum', 'sheet': 'Sales', 'group_by': ['Region'], 'sum_columns': ['Amount'], 'output_sheet': 'Totals'},
        {'op': 'formula', 'sheet': 'Sales', 'target_column': 'Revenue', 'formula': '=B{row}*C{row}'},
        {'op': 'format', 'sheet': 'Sales', 'column': 'Revenue', 'number_format': '0.00'},
    ])
    assert result['actions_count'] == 3
    book = load_workbook(result['output_path'])
    actual = {row[0]: row[1] for row in list(book['Totals'].values)[1:]}
    expected = {'East': sum(range(2, 1001, 2)), 'West': sum(range(1, 1001, 2))}
    assert actual == expected
    assert book['Sales']['A1'].font.bold
    assert book['Sales']['A1'].font.color.rgb == '00FF0000'
    assert book['Sales']['B501'].number_format == '#,##0.00'
    assert book['Sales']['D2'].value == '=B2*C2'
    book.close()
    assert hashlib.sha256((tmp_path / 'source.xlsx').read_bytes()).hexdigest() == before


def test_sort_filter_keep_cell_styles(tmp_path):
    make_workbook(tmp_path, 10)
    documents = WorkspaceDocuments(tmp_path)
    result = documents.transform_workbook('source.xlsx', 'filtered.xlsx', [
        {'op': 'filter', 'sheet': 'Sales', 'column': 'Amount', 'operator': 'gt', 'value': 5},
        {'op': 'sort', 'sheet': 'Sales', 'column': 'Amount', 'descending': True},
    ])
    book = load_workbook(result['output_path'])
    assert [row[1] for row in list(book['Sales'].values)[1:]] == [10, 9, 8, 7, 6]
    assert book['Sales']['B2'].number_format == '#,##0.00'
    book.close()


@pytest.mark.parametrize('source,output', [('../source.xlsx', 'out.xlsx'), ('source.xlsx', '../out.xlsx'), ('source.xlsx', 'source.xlsx'), ('source.xlsm', 'out.xlsx'), ('C:\\source.xlsx', 'out.xlsx')])
def test_paths_formats_and_overwrite_rejected(tmp_path, source, output):
    make_workbook(tmp_path, 2)
    with pytest.raises((ValueError, FileExistsError)):
        WorkspaceDocuments(tmp_path).transform_workbook(source, output, [])


def test_failed_operation_produces_no_partial_output(tmp_path):
    make_workbook(tmp_path, 2)
    with pytest.raises(ValueError):
        WorkspaceDocuments(tmp_path).transform_workbook('source.xlsx', 'out.xlsx', [
            {'op': 'formula', 'sheet': 'Sales', 'target_column': 'Calculated', 'formula': '=B{row}*2'},
            {'op': 'group_sum', 'sheet': 'Sales', 'group_by': ['Region'], 'sum_columns': ['Calculated'], 'output_sheet': 'Totals'},
        ])
    assert sorted(p.name for p in tmp_path.iterdir()) == ['source.xlsx']


def test_dangerous_formulas_and_row_bounds_rejected(tmp_path):
    make_workbook(tmp_path, 5)
    with pytest.raises(ValueError, match='limit'):
        WorkspaceDocuments(tmp_path, max_rows=4).inspect_workbook('source.xlsx')
    with pytest.raises(ValueError):
        WorkspaceDocuments(tmp_path).transform_workbook('source.xlsx', 'out.xlsx', [
            {'op': 'formula', 'sheet': 'Sales', 'target_column': 'Bad', 'formula': '=WEBSERVICE("https://example.com")'},
        ])


def test_formula_header_injection_rejected(tmp_path):
    make_workbook(tmp_path, 2)
    with pytest.raises(ValueError):
        WorkspaceDocuments(tmp_path).transform_workbook('source.xlsx', 'out.xlsx', [
            {'op': 'formula', 'sheet': 'Sales', 'target_column': '=WEBSERVICE("https://example.com")', 'formula': '=B{row}*2'},
        ])


def test_existing_output_untouched_and_empty_workbook_inspection(tmp_path):
    make_workbook(tmp_path, 0)
    (tmp_path / 'out.xlsx').write_bytes(b'existing output')
    with pytest.raises(FileExistsError):
        WorkspaceDocuments(tmp_path).transform_workbook('source.xlsx', 'out.xlsx', [])
    assert (tmp_path / 'out.xlsx').read_bytes() == b'existing output'
    assert WorkspaceDocuments(tmp_path).inspect_workbook('source.xlsx')['sheets'][0]['rows'] == 0


def test_sort_rejects_cross_sheet_formula_references(tmp_path):
    make_workbook(tmp_path, 2)
    book = load_workbook(tmp_path / 'source.xlsx')
    sheet = book.create_sheet('Links')
    sheet.append(['Amount'])
    sheet.append(['=Sales!B2'])
    book.save(tmp_path / 'source.xlsx')
    book.close()
    with pytest.raises(ValueError, match='formula'):
        WorkspaceDocuments(tmp_path).transform_workbook('source.xlsx', 'out.xlsx', [
            {'op': 'sort', 'sheet': 'Sales', 'column': 'Amount'},
        ])


@pytest.mark.parametrize('formula', ['=ZZZ{row}*2', '=B9999999*2', '=B{row}**2'])
def test_formula_references_and_grammar_bounded(tmp_path, formula):
    make_workbook(tmp_path, 2)
    with pytest.raises(ValueError):
        WorkspaceDocuments(tmp_path).transform_workbook('source.xlsx', 'out.xlsx', [
            {'op': 'formula', 'sheet': 'Sales', 'target_column': 'Derived', 'formula': formula},
        ])
