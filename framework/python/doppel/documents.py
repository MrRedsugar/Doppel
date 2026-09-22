"""Structured transformations of explicitly authorized local XLSX copies."""

from copy import copy
from datetime import date, datetime, time
import hashlib
import math
import os
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from .paths import contained_path


class WorkspaceDocuments:
    def __init__(self, root: str | Path, *, max_file_bytes: int = 20 * 1024 * 1024,
                 max_rows: int = 50000, max_columns: int = 256, max_cells: int = 1000000):
        self.root = Path(root).resolve(strict=True)
        self.max_file_bytes = max_file_bytes
        self.max_rows = max_rows
        self.max_columns = max_columns
        self.max_cells = max_cells
        if min(max_file_bytes, max_rows, max_columns, max_cells) <= 0:
            raise ValueError('Document limits must be positive')

    def _path(self, name: str) -> Path:
        path = contained_path(self.root, name)
        if path.suffix.lower() != '.xlsx':
            raise ValueError('Only .xlsx is supported; macros and .xlsm are unsupported')
        return path

    def _load(self, path: Path, *, transform: bool = False):
        if path.stat().st_size > self.max_file_bytes:
            raise ValueError('Workbook file byte limit exceeded')
        try:
            with ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > 10000 or sum(m.file_size for m in members) > self.max_file_bytes * 10:
                    raise ValueError('Workbook expanded size limit exceeded')
                names = [m.filename.lower() for m in members]
                if any('vbaproject' in n or n.startswith('xl/externallinks/') for n in names):
                    raise ValueError('Macros and external workbook links are unsupported')
                if transform and any(n.startswith(('xl/drawings/', 'xl/pivottables/', 'xl/slicers/')) for n in names):
                    raise ValueError('Drawing, pivot and slicer preservation is unsupported')
            book = load_workbook(path, data_only=False, keep_links=False)
        except BadZipFile as exc:
            raise ValueError('Invalid XLSX archive') from exc
        try:
            cells = 0
            for sheet in book:
                cells += sheet.max_row * sheet.max_column
                if sheet.max_row > self.max_rows + 1 or sheet.max_column > self.max_columns or cells > self.max_cells:
                    raise ValueError('Workbook row, column or cell limit exceeded')
            if len(book.sheetnames) > 100:
                raise ValueError('Workbook sheet count limit exceeded')
            return book
        except BaseException:
            book.close()
            raise

    @staticmethod
    def _value(value):
        if isinstance(value, (datetime, date, time)):
            return value.isoformat()
        if isinstance(value, str):
            return value[:256]
        return value

    def _summary(self, book):
        sheets = []
        for sheet in book:
            stats = {}
            for column, cells in enumerate(sheet.iter_cols(min_row=2), 1):
                numeric = [c.value for c in cells if c.data_type != 'f' and isinstance(c.value, (int, float)) and not isinstance(c.value, bool) and math.isfinite(c.value)]
                stats[get_column_letter(column)] = {'numeric_count': len(numeric), 'sum': sum(numeric), 'formula_count': sum(c.data_type == 'f' for c in cells)}
            sheets.append({'name': sheet.title, 'rows': max(0, sheet.max_row - 1), 'columns': sheet.max_column,
                           'headers': [self._value(c.value) for c in sheet[1]],
                           'samples': [[self._value(c.value) for c in row] for row in sheet.iter_rows(min_row=2, max_row=min(sheet.max_row, 6))] if sheet.max_row > 1 else [],
                           'statistics': stats})
        return {'sheets': sheets, 'formula_recalculation': 'not_performed', 'trusted': False}

    def inspect_workbook(self, name: str) -> dict:
        book = self._load(self._path(name))
        try:
            return self._summary(book)
        finally:
            book.close()

    @staticmethod
    def _columns(sheet):
        headers = [cell.value for cell in sheet[1]]
        if not headers or any(not isinstance(h, str) or not h for h in headers) or len(set(headers)) != len(headers):
            raise ValueError('Operations require unique nonempty text headers in row one')
        return {header: i + 1 for i, header in enumerate(headers)}

    @staticmethod
    def _column(columns, name):
        if name not in columns:
            raise ValueError(f'Unknown column: {name}')
        return columns[name]

    @staticmethod
    def _rewrite_rows(sheet, rows):
        snapshots = [[(c.value, copy(c._style), copy(c.comment), copy(c.hyperlink)) for c in row] for row in rows]
        dimensions = [copy(sheet.row_dimensions[row[0].row]) for row in rows]
        sheet.delete_rows(2, max(0, sheet.max_row - 1))
        for key in list(sheet.row_dimensions):
            if key >= 2:
                del sheet.row_dimensions[key]
        for row_number, (row, dimension) in enumerate(zip(snapshots, dimensions), 2):
            dimension.index = row_number
            sheet.row_dimensions[row_number] = dimension
            for col, (value, style, comment, hyperlink) in enumerate(row, 1):
                cell = sheet.cell(row_number, col, value)
                cell._style, cell.comment, cell.hyperlink = style, comment, hyperlink

    def _apply(self, book, operation):
        contracts = {
            'sort': ({'op', 'sheet', 'column'}, {'descending'}),
            'filter': ({'op', 'sheet', 'column', 'operator', 'value'}, set()),
            'formula': ({'op', 'sheet', 'target_column', 'formula'}, set()),
            'group_sum': ({'op', 'sheet', 'group_by', 'sum_columns', 'output_sheet'}, set()),
            'format': ({'op', 'sheet', 'column'}, {'number_format', 'bold'}),
        }
        if not isinstance(operation, dict) or operation.get('op') not in contracts:
            raise ValueError('Unknown structured workbook operation')
        required, optional = contracts[operation['op']]
        if not required <= operation.keys() or operation.keys() - required - optional:
            raise ValueError('Missing or unknown operation fields')
        if operation['sheet'] not in book.sheetnames:
            raise ValueError('Unknown sheet')
        sheet = book[operation['sheet']]
        columns = self._columns(sheet)
        kind = operation['op']
        if kind in ('sort', 'filter'):
            if sheet.merged_cells or sheet.tables or len(sheet.conditional_formatting) or sheet.data_validations.count:
                raise ValueError('Row reordering with merged cells, tables or rules is unsupported')
            if any(c.data_type == 'f' for page in book for row in page for c in row):
                raise ValueError('Sorting/filtering formula rows requires recalculation support')
            col = self._column(columns, operation['column']) - 1
            rows = list(sheet.iter_rows(min_row=2)) if sheet.max_row > 1 else []
            if kind == 'sort':
                descending = operation.get('descending', False)
                if not isinstance(descending, bool):
                    raise ValueError('descending must be boolean')
                try:
                    rows.sort(key=lambda row: (row[col].value is None, row[col].value), reverse=descending)
                except TypeError as exc:
                    raise ValueError('Sort column contains incompatible types') from exc
            else:
                comparator, value = operation['operator'], operation['value']
                comparisons = {'eq': lambda a: a == value, 'ne': lambda a: a != value,
                               'gt': lambda a: a > value, 'gte': lambda a: a >= value,
                               'lt': lambda a: a < value, 'lte': lambda a: a <= value,
                               'contains': lambda a: isinstance(a, str) and isinstance(value, str) and value in a}
                if comparator not in comparisons:
                    raise ValueError('Unknown filter operator')
                try:
                    rows = [row for row in rows if comparisons[comparator](row[col].value)]
                except TypeError as exc:
                    raise ValueError('Filter value type is incompatible with column') from exc
            self._rewrite_rows(sheet, rows)
        elif kind == 'formula':
            header, formula = operation['target_column'], operation['formula']
            if not isinstance(header, str) or not header or len(header) > 256 or header.startswith('=') or header in columns or len(columns) >= self.max_columns:
                raise ValueError('Formula requires a new bounded column')
            if not isinstance(formula, str) or len(formula) > 512 or not formula.startswith('='):
                raise ValueError('Invalid formula')
            expression = formula[1:].replace('{row}', '2')
            tokens = re.findall(r'[A-Z]{1,3}[1-9][0-9]*|(?:[0-9]+(?:\.[0-9]+)?)|[+*/() -]', expression)
            if ''.join(tokens) != expression or not expression.strip():
                raise ValueError('Only local cell references, numbers and arithmetic formulas are supported')
            for column, row in re.findall(r'([A-Z]{1,3})([1-9][0-9]*)', expression):
                if column_index_from_string(column) > len(columns) or int(row) > sheet.max_row:
                    raise ValueError('Formula references must stay within existing sheet data')
            # The grammar is deliberately smaller than Excel: no functions, links or strings.
            import ast
            try:
                tree = ast.parse(re.sub(r'[A-Z]{1,3}[1-9][0-9]*', '1', expression), mode='eval')
            except SyntaxError as exc:
                raise ValueError('Invalid arithmetic formula') from exc
            if any(not isinstance(node, (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Add, ast.Sub, ast.Mult, ast.Div, ast.UAdd, ast.USub)) for node in ast.walk(tree)):
                raise ValueError('Unsupported arithmetic formula')
            target = len(columns) + 1
            sheet.cell(1, target, header)
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, target, formula.replace('{row}', str(row)))
            book.calculation.fullCalcOnLoad = True
        elif kind == 'group_sum':
            groups, sums = operation['group_by'], operation['sum_columns']
            if not isinstance(groups, list) or not isinstance(sums, list) or not groups or not sums or any(not isinstance(c, str) for c in groups + sums) or len(set(groups + sums)) != len(groups + sums):
                raise ValueError('Aggregation requires distinct grouping and numeric columns')
            group_indexes = [self._column(columns, name) - 1 for name in groups]
            sum_indexes = [self._column(columns, name) - 1 for name in sums]
            output_sheet = operation['output_sheet']
            if not isinstance(output_sheet, str) or not output_sheet or len(output_sheet) > 31 or re.search(r'[\\/*?:\[\]]', output_sheet) or output_sheet.lower() in [s.lower() for s in book.sheetnames]:
                raise ValueError('Aggregation requires a new valid output sheet')
            aggregated = {}
            for row in sheet.iter_rows(min_row=2):
                if any(row[i].data_type == 'f' for i in group_indexes + sum_indexes):
                    raise ValueError('Cannot aggregate formulas without trusted recalculated values')
                key = tuple(row[i].value for i in group_indexes)
                total = aggregated.setdefault(key, [0] * len(sums))
                for n, i in enumerate(sum_indexes):
                    value = row[i].value
                    if value is None:
                        continue
                    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                        raise ValueError('Aggregation sum columns must contain finite numbers')
                    total[n] += value
            target = book.create_sheet(output_sheet)
            target.append(groups + sums)
            for key, values in aggregated.items():
                target.append(list(key) + values)
        elif kind == 'format':
            col = self._column(columns, operation['column'])
            if 'number_format' not in operation and 'bold' not in operation:
                raise ValueError('Format requires number_format or bold')
            if 'number_format' in operation and (not isinstance(operation['number_format'], str) or len(operation['number_format']) > 128):
                raise ValueError('Invalid number format')
            if 'bold' in operation and not isinstance(operation['bold'], bool):
                raise ValueError('bold must be boolean')
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, col)
                if 'number_format' in operation:
                    cell.number_format = operation['number_format']
                if 'bold' in operation:
                    font = copy(cell.font)
                    font.bold = operation['bold']
                    cell.font = font

    def transform_workbook(self, source: str, output: str, operations: list[dict]) -> dict:
        source_path, output_path = self._path(source), self._path(output)
        if output_path.exists():
            raise FileExistsError('Output already exists; select a new version name')
        if not output_path.parent.is_dir():
            raise ValueError('Output parent must already exist in the workspace')
        if not isinstance(operations, list) or len(operations) > 20:
            raise ValueError('Operations must be an array of at most twenty objects')
        before = hashlib.sha256(source_path.read_bytes()).digest() if source_path.stat().st_size <= self.max_file_bytes else None
        book = self._load(source_path, transform=True)
        staging, reserved = None, False
        try:
            for operation in operations:
                self._apply(book, operation)
            handle, staging = tempfile.mkstemp(prefix='.doppel-', suffix='.xlsx', dir=output_path.parent)
            os.close(handle)
            book.save(staging)
            checked = self._load(Path(staging))
            try:
                summary = self._summary(checked)
            finally:
                checked.close()
            if hashlib.sha256(source_path.read_bytes()).digest() != before:
                raise ValueError('Source changed during transformation; import a fresh authorized copy')
            # Reserve with O_EXCL so another writer cannot be silently overwritten.
            self._path(output)
            descriptor = os.open(output_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
            os.close(descriptor)
            reserved = True
            os.replace(staging, output_path)
            staging, reserved = None, False
            return {'output_path': str(output_path), 'actions_count': len(operations), 'summary': summary,
                    'source_unchanged': True, 'reopened': True}
        finally:
            book.close()
            if staging is not None:
                Path(staging).unlink(missing_ok=True)
            if reserved:
                output_path.unlink(missing_ok=True)
